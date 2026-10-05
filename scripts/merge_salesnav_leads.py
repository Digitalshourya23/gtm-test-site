"""Merge salesnav/leads.csv into AWS_Decision_Makers_Contacts.xlsx as a 'Sales Nav Leads' sheet.

Usage: python3 scripts/merge_salesnav_leads.py
Requires: pip install openpyxl

Optional enrichment: if salesnav/enriched.csv exists (an export from Apollo, Lusha or ContactOut),
its email/phone columns are matched to leads by salesnav_lead_url, or by name + company.
"""
import csv
import os
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "AWS_Decision_Makers_Contacts.xlsx")
LEADS = os.path.join(ROOT, "salesnav", "leads.csv")
ENRICHED = os.path.join(ROOT, "salesnav", "enriched.csv")

TECH = re.compile(r"\b(cto|cio|cdo|cdio|chief (technology|information|digital|data|product)|"
                  r"engineering|technology|infrastructure|cloud|devops|platform|it\b|information technology|architect|sre)", re.I)
TOP = re.compile(r"\b(chief|cxo|cto|cio|ceo|founder|co-founder|(?<!vice )president|managing director|md\b)", re.I)


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def rank(title):
    t = title or ""
    if TOP.search(t) and TECH.search(t):
        return "1 - Tech C-level"
    if TECH.search(t) and re.search(r"\b(vp|vice president|head|svp|evp|avp)\b", t, re.I):
        return "2 - Tech VP/Head"
    if TECH.search(t):
        return "3 - Tech Director/Lead"
    if TOP.search(t):
        return "4 - Business C-level"
    return "5 - Other senior"


def load_enriched():
    by_url, by_name = {}, {}
    if not os.path.exists(ENRICHED):
        return by_url, by_name
    with open(ENRICHED, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            low = {k.lower().strip(): (v or "").strip() for k, v in r.items() if k}

            def pick(*keys):
                for k in keys:
                    for col, v in low.items():
                        if k in col and v:
                            return v
                return ""

            rec = {
                "email": pick("work email", "email"),
                "phone1": pick("mobile", "direct", "phone"),
                "phone2": pick("corporate phone", "company phone", "other phone"),
            }
            url = pick("salesnav", "linkedin")
            first, last = low.get("first name", ""), low.get("last name", "")
            name = f"{first} {last}".strip() if first or last else pick("full name", "name")
            if url:
                by_url[url.split("?")[0]] = rec
            by_name[norm(name) + "|" + norm(pick("company", "organization"))] = rec
    return by_url, by_name


def main():
    if not os.path.exists(LEADS):
        raise SystemExit(f"No leads file yet: {LEADS}")
    with open(LEADS, newline="", encoding="utf-8") as f:
        leads = list(csv.DictReader(f))
    by_url, by_name = load_enriched()

    wb = load_workbook(XLSX)
    if "Sales Nav Leads" in wb.sheetnames:
        del wb["Sales Nav Leads"]
    ws = wb.create_sheet("Sales Nav Leads", 1)
    header = ["List", "#", "Company", "Name", "Title", "Priority", "Location", "Time in role",
              "Degree", "Sales Nav lead URL", "Email", "Phone - Primary", "Phone - Secondary", "Collected at"]
    ws.append(header)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True)

    seen = set()
    leads.sort(key=lambda r: (r["list"], int(r["num"] or 0), rank(r["title"])))
    for r in leads:
        key = r["salesnav_lead_url"].split("?")[0] or norm(r["name"]) + norm(r["company"])
        if key in seen:
            continue
        seen.add(key)
        e = by_url.get(r["salesnav_lead_url"].split("?")[0]) or \
            by_name.get(norm(r["name"]) + "|" + norm(r["company_shown"] or r["company"])) or {}
        url = r["salesnav_lead_url"]
        if url.startswith("/"):
            url = "https://www.linkedin.com" + url
        ws.append([r["list"], int(r["num"] or 0), r["company"], r["name"], r["title"], rank(r["title"]),
                   r["location"], r["time_in_role"], r["degree"], url,
                   e.get("email", ""), e.get("phone1", ""), e.get("phone2", ""), r["collected_at"]])
        cell = ws.cell(ws.max_row, 10)
        if url:
            cell.hyperlink = url
            cell.font = Font(color="0563C1", underline="single")

    for col, w in zip("ABCDEFGHIJKLMN", [12, 5, 32, 26, 44, 20, 24, 16, 8, 45, 30, 18, 18, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(XLSX)
    print(f"Merged {len(seen)} unique leads into 'Sales Nav Leads' ({XLSX})")


if __name__ == "__main__":
    main()
